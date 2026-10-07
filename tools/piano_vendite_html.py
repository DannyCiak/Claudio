"""Crea la versione HTML interattiva del piano a partire dall'Excel generato da piano_vendite_portafoglio.py.

Uso: python3 tools/piano_vendite_html.py Piano_Vendite_Full_Immersion.xlsx Piano_Vendite_Interattivo.html
"""
import datetime as dt
import json
import pathlib
import sys

import openpyxl

K = {'ID': 'id', 'Cliente': 'n', 'Nome per saluto': 's', 'Tipo': 't', 'Sesso': 'sx', 'Età': 'e', 'Professione': 'pr',
     'Comune (Prov)': 'co', 'Cellulare': 'ph', 'Email': 'em', 'Lei / tu': 'lt', 'Privacy marketing': 'pv',
     'Firma OTP': 'otp', 'Quest. IBIPs': 'ib', 'Cliente dal': 'cd', 'N. polizze': 'np', 'Premi annui €': 'pa',
     'Prossima quietanza': 'sq', 'Portafoglio attuale': 'pf', 'Moduli salute attuali': 'ms',
     'Iniziative Allianz assegnate': 'ini', 'ALERT': 'al', 'Score': 'sc', 'Classe': 'cl', 'Prodotto 1 (focus)': 'p1',
     'Perché (dai dati)': 'why', 'Prodotto 2': 'p2', 'Prodotto 3': 'p3', 'Data 1° contatto': 'd', 'Ora': 'h',
     'Sequenza di contatto': 'seq', 'Tipo di hook': 'hk', 'SCRIPT TELEFONATA': 'scr', 'WhatsApp 1 – preavviso': 'wa1',
     'WhatsApp 2 – se non risponde': 'wa2', 'Email (3° tentativo)': 'mail',
     'Domande di analisi (appuntamento 1)': 'dq', 'Obiezione probabile → risposta': 'ob',
     'Proposta da presentare (Good/Better/Best)': 'prop', 'Documenti da portare': 'docs',
     'Luogo appuntamento': 'lu', 'Appuntamento: 2 opzioni da proporre': 'app', 'Firma': 'fi', 'Compliance': 'cmp'}
G = {'Gap TCM': 'TCM', 'Gap Salute': 'Salute', 'Gap LTC': 'Non autosuff.', 'Gap Previdenza': 'Previdenza',
     'Gap Risparmio/Inv.': 'Risparmio', 'Gap Casa': 'Casa', 'Gap Auto': 'Auto', 'Gap Impresa/CAT NAT': 'Impresa/CAT NAT'}

src, out = sys.argv[1], sys.argv[2]
ws = openpyxl.load_workbook(src, data_only=True)['Schede Clienti']
hdr = [c.value for c in ws[1]]
rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r, gaps = {}, []
    for h, v in zip(hdr, row):
        if isinstance(v, (dt.datetime, dt.date)):
            v = v.strftime('%Y-%m-%d')
        if h in K:
            r[K[h]] = v if v not in (None, '—') else ''
        elif h in G and v == 'Sì':
            gaps.append(G[h])
    r['g'] = gaps
    r['wn'] = ''.join(ch for ch in str(r['ph']) if ch.isdigit())
    rows.append(r)
rows.sort(key=lambda r: (r['d'], r['h'], r['id']))
data = json.dumps(rows, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
tpl = (pathlib.Path(__file__).parent / 'piano_vendite_template.html').read_text()
pathlib.Path(out).write_text(tpl.replace('__DATA__', data))
print(f'{len(rows)} clienti → {out}')
